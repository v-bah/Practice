# Импорт библиотек
from flask import Flask, render_template, request, url_for
from flask_sqlalchemy import SQLAlchemy
from docx import Document
import os
# Запуск сервера и связь с базой данных на будущее
app = Flask(__name__)
app.config['SQLALCHEMY_DATABASE_URI'] = 'sqlite:///phoenix.db'
Users_db = SQLAlchemy(app)
document = Document('/static/contract_example.docx')


class Users(Users_db.Model):
    id = Users_db.Column(Users_db.Integer, primary_key=True)
    name = Users_db.Column(Users_db.String(1500), nullable=False)
    birth_date = Users_db.Column(Users_db.Integer, nullable=False)
    INN = Users_db.Column(Users_db.Integer, nullable=False)
    passport_num = Users_db.Column(Users_db.Integer, primarnullabley_key=False)
    passport_place = Users_db.Column(Users_db.Text, nullable=False)
    passport_date = Users_db.Column(Users_db.Integer, nullable=False)
    passport_code = Users_db.Column(Users_db.Integer, nullable=False)
    address = Users_db.Column(Users_db.Text, nullable=False)
    bank_account = Users_db.Column(Users_db.Integer, nullable=False)
    bank_name = Users_db.Column(Users_db.String(50), nullable=False)
    bank_details = Users_db.Column(Users_db.Text, nullable=False)
    email = Users_db.Columb(Users_db.String(100), nullable = False)

    def __repr__(self):
        return '<UsersContract %r>' % self.id


# # Создание словаря. Заполнен тестовыми переменными
# person_account = {
#     'name': 'Test Object 2-gou',
#     'nickname': 'Test',
#     'vk_link': 'vk.com/poisonowl',
#     'INN': '119112475684',
#     'bank_details': 'БИК: 044525225, КПП: 773601001, ИНН: 7707083893',
#     'bank_name': 'Sberbank',
#     'password': '1234',
#     'email': 'test_object@gmail.com',
#     'phone_number': '+79777777777'
# }
#
#
# # Запуск странички с отображением данных личного кабинета
# @app.route('/cabinet', methods=['GET'])
# def cabinet():
#     url_sequence = 'test'
#     person_account[url_sequence] = person_account
#     return render_template('cabinet_test.html', person_account = person_account, url_sequence = url_sequence)
#
#
# # Запуск странички с изменением существующих данных.
# @app.route('/info_change', methods=['POST','GET'])
# def info_change():
#     if request.method=="POST":
#         name = request.form.get('name')
#         nickname = request.form.get('nickname')
#         vk_link = request.form.get('vk_link')
#         INN = request.form.get('inn')
#         bank_details = request.form.get('bank_details')
#         bank_name = request.form.get('bank_name')
#         current_password = person_account['password']
#         new_pass = request.form.get('new_pass')
#         email = request.form.get('email')
#         phone_number = request.form.get('phone_number')
#         if current_password == request.form.get('old_pass'):  # Проверка на соответствие старого пароля, полученного с формы
#             if new_pass == request.form.get('repeat_pass'):  # Проверка на новый пароль, чтобы пользователь не ошибся
#                 new_pass = new_pass
#                 # Очищение и заполнение словаря данными
#                 person_account.clear()
#                 person_account['name'] = name
#                 person_account['nickname'] = nickname
#                 person_account['vk_link'] = vk_link
#                 person_account['INN'] = INN
#                 person_account['bank_details'] = bank_details
#                 person_account['bank_name'] = bank_name
#                 person_account['password'] = new_pass
#                 person_account['email'] = email
#                 person_account['phone_number'] = phone_number
#                 return (person_account)
#             else:
#                 return "error"
#         else:
#             return "error"
#     return  render_template('info_change.html', person_account=person_account)
#
#
# @app.route('/login', methods = ['POST','GET'])
# def login():
#     return render_template('login.html')
#
#
# @app.route('/register', methods = ['POST', 'GET'])
# def register():
#     return render_template('register.html')
#

@app.route('/create_contract', methods=['POST', 'GET'], User=Users)
def contract():
    name = ''
    id_sel = 1
    if request.method == 'POST':
        name = request.form.get('name')  # запрос к данным формы
        id_sel = request.form.get('id')
        table = document.add_table(rows=1, cols=1)
        cell = table.cell(0, 1)
        birth_date = Users.birth_date
        INN = Users.INN
        passport_num = Users.passport_num
        passport_place = Users.passport_place
        passport_date = Users.passport_date
        passport_code = Users.passport_code
        address = Users.address
        bank_account = Users.bank_account
        bank_name = Users.bank_name
        bank_details = Users.bank_details
        email = Users.email
        cell.text = "Заказчик " \
                    "Индивидуальный предприниматель" \
                    "Нечитайло Фёдор Константинович" \
                    "ИНН 616616300580 ОГРН 318619600017594" \
                    "Адрес: 344065, Ростовская обл., г. Ростов-на-Дону, ул. Вятская, д. 63/1, кв. 77" \
                    "р/с 40802810000000405802 в АО «Тинькофф Банк»" \
                    "кор/сч 30101810145250000974" \
                    "БИК 044525974" \
                    "fedorcomixvideo@gmail.com"
        cell = table.call(0, 0)
        cell.text = "Подрядчик " + name + birth_date + INN + passport_num + passport_place + passport_date + passport_code \
                    + address + bank_account + bank_name + bank_details + email
    return render_template('create_contract.html', name=name, id=id_sel)


if __name__ == "__main__":
    app.run(debug=True)
